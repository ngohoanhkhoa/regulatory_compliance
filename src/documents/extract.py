"""Text extraction for uploaded user documents.

Supported formats: PDF (``pypdf``), DOCX (``python-docx``), and plain text
(``.txt``/``.md``/``.markdown``/``.csv``). Legacy binary ``.doc`` is not
supported (requires LibreOffice); callers get a clear error.
"""

from __future__ import annotations

import io
from pathlib import Path

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(
    {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv"}
)


class UnsupportedDocument(RuntimeError):
    """Raised for a file extension we do not extract."""


class DocumentExtractionError(RuntimeError):
    """Raised when a supported file cannot be parsed."""


def extension_of(filename: str | None) -> str:
    return Path(filename or "").suffix.lower()


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from ``data``. Raises on unsupported/corrupt input."""
    ext = extension_of(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(e.lstrip(".") for e in SUPPORTED_EXTENSIONS))
        raise UnsupportedDocument(
            f"Unsupported file type '{ext or filename}'. Allowed: {allowed}."
        )
    try:
        if ext == ".pdf":
            return _extract_pdf(data)
        if ext == ".docx":
            return _extract_docx(data)
        return data.decode("utf-8", errors="replace")
    except UnsupportedDocument:
        raise
    except Exception as exc:  # library-specific parse failures
        raise DocumentExtractionError(f"Could not read {filename}: {exc}") from exc


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
